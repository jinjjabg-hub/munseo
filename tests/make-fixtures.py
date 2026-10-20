#!/usr/bin/env python3
"""테스트용 문서를 만든다 (tests/.cache/fixtures). 사용: python3 tests/make-fixtures.py
- 복잡한 .docx, 표·그림 있는 .pptx, 시트 여러 개 + 6만 행 .xlsx, .xls/.ods/.csv(CP949)
- 100쪽 .pdf, 암호 걸린 .pdf/.docx, 암호 표시 .hwp(실제 .hwp의 암호 플래그만 켠 것), .hwpx, 깨진 파일
"""
import os, sys, io, re, struct, shutil, subprocess, random, zipfile
OUT = os.path.join(os.path.dirname(__file__), '.cache', 'fixtures')
os.makedirs(OUT, exist_ok=True)
REAL_HWP = os.environ.get('REAL_HWP_DIR', '/tmp/claude-0/real-hwp')
random.seed(7)

from PIL import Image, ImageDraw

def png(path, w=640, h=360, color=(40, 110, 220), label='그림 IMG'):
    im = Image.new('RGB', (w, h), color)
    d = ImageDraw.Draw(im)
    for i in range(0, w, 40): d.line([(i, 0), (i, h)], fill=(255, 255, 255), width=1)
    for i in range(0, h, 40): d.line([(0, i), (w, i)], fill=(255, 255, 255), width=1)
    d.rectangle([20, 20, w - 20, h - 20], outline=(255, 255, 0), width=6)
    d.text((40, 40), label, fill=(255, 255, 255))
    im.save(path)

# ---------- docx (서식 복잡) ----------
def make_docx():
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT
    img = os.path.join(OUT, '_img.png'); png(img)
    d = Document()
    sec = d.sections[0]
    sec.header.paragraphs[0].text = '머리말: 월간 업무 보고서 (Header)'
    sec.footer.paragraphs[0].text = '꼬리말: 사내용 — Footer'
    d.add_heading('월간 업무 보고서 Monthly Report', 0)
    p = d.add_paragraph('이 문서는 ')
    r = p.add_run('굵은 글씨'); r.bold = True
    p.add_run(', ')
    r = p.add_run('기울임'); r.italic = True
    p.add_run(', ')
    r = p.add_run('밑줄'); r.underline = True
    p.add_run(', ')
    r = p.add_run('빨간색 큰 글씨'); r.font.color.rgb = RGBColor(220, 30, 30); r.font.size = Pt(16)
    p.add_run(' 서식을 포함합니다. 한글과 English가 섞여 있고 숫자 1,234,567도 있습니다.')
    d.add_heading('1. 개요', 1)
    d.add_paragraph('첫 번째 불릿 항목', style='List Bullet')
    d.add_paragraph('두 번째 불릿 항목', style='List Bullet')
    d.add_paragraph('번호 항목 하나', style='List Number')
    d.add_paragraph('번호 항목 둘', style='List Number')
    d.add_heading('2. 실적 표', 1)
    t = d.add_table(rows=4, cols=4); t.style = 'Table Grid'
    hdr = ['구분', '1월', '2월', '3월']
    for i, h in enumerate(hdr): t.cell(0, i).text = h
    rows = [('매출', '1,200', '1,350', '1,500'), ('비용', '800', '820', '900')]
    for ri, row in enumerate(rows, 1):
        for ci, v in enumerate(row): t.cell(ri, ci).text = v
    a = t.cell(3, 0).merge(t.cell(3, 3)); a.text = '병합된 셀: 합계 및 비고'
    d.add_heading('3. 그림', 1)
    d.add_picture(img, width=Cm(10))
    cap = d.add_paragraph('그림 1. 샘플 이미지'); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i in range(40):
        d.add_paragraph(f'본문 문단 {i+1}: 가나다라마바사아자차카타파하 The quick brown fox jumps over the lazy dog. ' * 2)
    d.add_page_break()
    d.add_heading('4. 마지막 쪽', 1)
    d.add_paragraph('문서의 마지막 문단입니다. END-OF-DOCX')
    d.save(os.path.join(OUT, 'complex.docx'))

# ---------- pptx (표 + 그림) ----------
def make_pptx():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    img = os.path.join(OUT, '_img2.png'); png(img, color=(200, 80, 60), label='PPT 그림')
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = '2025 상반기 사업 보고'
    s.placeholders[1].text = '내 문서 뷰어 시험용 · Quarterly Review'
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = '핵심 요약'
    tf = s.placeholders[1].text_frame
    tf.text = '매출 12% 증가'
    for line in ['신규 고객 340곳 확보', '해외 시장 2개국 진출', 'Churn rate 하락']:
        p = tf.add_paragraph(); p.text = line
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = '분기별 실적 표'
    rows, cols = 4, 4
    tbl = s.shapes.add_table(rows, cols, Inches(1), Inches(2), Inches(11), Inches(3)).table
    data = [['분기', '매출', '비용', '이익'], ['1Q', '120', '80', '40'], ['2Q', '135', '85', '50'], ['3Q', '150', '90', '60']]
    for r in range(rows):
        for c in range(cols): tbl.cell(r, c).text = data[r][c]
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = '그림 슬라이드'
    s.shapes.add_picture(img, Inches(2), Inches(1.8), width=Inches(8))
    tb = s.shapes.add_textbox(Inches(2), Inches(6.6), Inches(9), Inches(0.6))
    tb.text_frame.text = '그림 설명 텍스트 END-OF-PPTX'
    prs.save(os.path.join(OUT, 'table-picture.pptx'))

# ---------- xlsx ----------
def make_xlsx():
    import openpyxl, datetime
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = '요약'
    ws.append(['항목', '금액', '날짜', '비고'])
    for i in range(1, 30):
        ws.append([f'품목 {i}', i * 1234.5, datetime.date(2025, 1, 1) + datetime.timedelta(days=i), '한글 비고 텍스트' if i % 2 else 'note'])
    ws.merge_cells('F1:H2'); ws['F1'] = '병합 제목'
    ws.column_dimensions['A'].width = 18
    ws2 = wb.create_sheet('대용량 60000행')
    ws2.append(['번호', '이름', '값1', '값2', '값3', '메모', '합계', '플래그'])
    for i in range(1, 60001):
        ws2.append([i, f'이름{i}', i * 3, i * 7 % 1000, round(i / 7, 3), f'row-{i}-메모', None, i % 2 == 0])
    ws3 = wb.create_sheet('Sheet3 영문')
    for r in range(1, 40):
        ws3.append([f'R{r}C{c}' for c in range(1, 30)])
    wb.save(os.path.join(OUT, 'multi-sheet.xlsx'))
    # 같은 내용을 다른 형식으로 (LibreOffice)
    small = openpyxl.Workbook(); w = small.active; w.title = '시트A'
    w.append(['이름', '점수']); w.append(['홍길동', 90]); w.append(['김철수', 85])
    w2 = small.create_sheet('시트B'); w2.append(['x', 'y']); w2.append([1, 2])
    small.save(os.path.join(OUT, '_small.xlsx'))
    for fmt in ('xls', 'ods'):
        subprocess.run(['soffice', '--headless', '--convert-to', fmt, '--outdir', OUT, os.path.join(OUT, '_small.xlsx')], capture_output=True)
    shutil.move(os.path.join(OUT, '_small.xls'), os.path.join(OUT, 'legacy.xls'))
    shutil.move(os.path.join(OUT, '_small.ods'), os.path.join(OUT, 'sheet.ods'))
    with open(os.path.join(OUT, 'korean-cp949.csv'), 'wb') as f:
        f.write('이름,점수,메모\r\n홍길동,90,"쉼표, 포함"\r\n김철수,85,한글\r\n'.encode('cp949'))

# ---------- pdf ----------
def make_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    img = os.path.join(OUT, '_img.png')
    text = ('The quick brown fox jumps over the lazy dog. Lorem ipsum dolor sit amet, consectetur adipiscing elit, '
            'sed do eiusmod tempor incididunt ut labore et dolore magna aliqua. ')
    def build(path, pages, **kw):
        c = canvas.Canvas(path, pagesize=A4, **kw)
        W, H = A4
        for p in range(1, pages + 1):
            c.setFont('Helvetica-Bold', 20); c.drawString(50, H - 60, f'Page {p} of {pages}')
            c.setFont('Helvetica', 11)
            y = H - 100
            for i in range(38):
                c.drawString(50, y, f'{i+1:02d}. ' + text[:95]); y -= 16
            if p % 5 == 0:
                c.drawImage(ImageReader(img), 50, 80, width=300, height=170)
            c.setStrokeColorRGB(.2, .4, .8)
            for k in range(30): c.line(380, 100 + k * 3, 540, 120 + k * 5)
            c.drawString(50, 40, f'END-PAGE-{p}')
            c.showPage()
        c.save()
    build(os.path.join(OUT, 'hundred-pages.pdf'), 100)
    # 스캔본처럼 쪽마다 큰 사진이 들어간 100쪽 PDF (파일이 큼)
    sc = canvas.Canvas(os.path.join(OUT, 'scanned-100.pdf'), pagesize=A4)
    W, H = A4
    for p in range(1, 101):
        small = Image.frombytes('RGB', (124, 175), os.urandom(124 * 175 * 3)).resize((1240, 1754), Image.BICUBIC)
        ip = os.path.join(OUT, '_scan.jpg'); small.save(ip, quality=72)
        sc.drawImage(ip, 0, 0, width=W, height=H)
        sc.setFont('Helvetica-Bold', 28); sc.setFillColorRGB(1, 1, 1); sc.drawString(60, H - 80, f'SCAN PAGE {p}')
        sc.showPage()
    sc.save()
    from reportlab.lib.pdfencrypt import StandardEncryption
    enc = StandardEncryption('open123', canPrint=1)
    build(os.path.join(OUT, 'encrypted.pdf'), 3, encrypt=enc)

def make_korean_pdf():
    # LibreOffice로 한글 글꼴이 들어간 PDF 만들기
    from docx import Document
    d = Document()
    d.add_paragraph('한글 PDF 시험: 가나다라마바사 아자차카타파하 대한민국 END-OF-KPDF')
    p = os.path.join(OUT, '_k.docx'); d.save(p)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', OUT, p], capture_output=True)
    shutil.move(os.path.join(OUT, '_k.pdf'), os.path.join(OUT, 'korean.pdf'))

# ---------- 암호/깨진 파일 ----------
def make_protected():
    import msoffcrypto
    # 암호 걸린 docx (실제 OOXML 암호화)
    with open(os.path.join(OUT, 'complex.docx'), 'rb') as f:
        office = msoffcrypto.OfficeFile(f)
        office.load_key(password='pw1234')
    # msoffcrypto는 복호화용이라 암호화는 encrypt()로
    with open(os.path.join(OUT, 'complex.docx'), 'rb') as f:
        office = msoffcrypto.OfficeFile(f)
        with open(os.path.join(OUT, 'encrypted.docx'), 'wb') as o:
            office.encrypt('pw1234', o)
    # 암호/배포용 표시 .hwp: 실제 .hwp의 FileHeader 속성 비트만 켠 것
    src = os.path.join(REAL_HWP, 'naver', 'fax.hwp')
    if os.path.exists(src):
        b = bytearray(open(src, 'rb').read())
        # FileHeader 서명이 파일 안에 여러 번 나올 수 있어 모두 고친다
        pos = [m.start() for m in re.finditer(re.escape(b'HWP Document File'), bytes(b))]
        for name, bit in (('encrypted.hwp', 0x2), ('distribution.hwp', 0x4)):
            c = bytearray(b)
            for i in pos:
                flags = struct.unpack_from('<I', c, i + 36)[0]
                struct.pack_into('<I', c, i + 36, flags | bit)
            open(os.path.join(OUT, name), 'wb').write(c)
    # 깨진 파일들
    data = open(os.path.join(OUT, 'complex.docx'), 'rb').read()
    open(os.path.join(OUT, 'truncated.docx'), 'wb').write(data[: len(data) // 3])
    open(os.path.join(OUT, 'garbage.hwp'), 'wb').write(os.urandom(5000))
    open(os.path.join(OUT, 'garbage.pdf'), 'wb').write(b'%PDF-1.4\n' + os.urandom(3000))
    open(os.path.join(OUT, 'garbage.xlsx'), 'wb').write(b'PK\x03\x04' + os.urandom(3000))
    with zipfile.ZipFile(os.path.join(OUT, 'archive.zip'), 'w') as z:
        z.writestr('readme.txt', 'just a zip')
    open(os.path.join(OUT, 'unknown.xyz'), 'wb').write(b'hello')

def make_hwpx():
    try:
        from hwpx import HwpxDocument
    except Exception as e:
        print('hwpx skip', e); return
    doc = HwpxDocument.new()
    doc.add_paragraph('안녕하세요. 한글 HWPX 테스트 문서입니다.')
    doc.add_paragraph('두 번째 문단 — 가나다라마바사 ABC 123')
    t = doc.add_table(3, 3)
    for r in range(3):
        for c in range(3): t.set_cell_text(r, c, f'칸{r+1}-{c+1}')
    t.merge_cells(0, 0, 0, 1)
    doc.add_paragraph('표 아래 문단입니다. END-OF-HWPX')
    doc.save_to_path(os.path.join(OUT, 'sample.hwpx'))

def copy_real_hwp():
    # 사용자가 test-files에 넣어둔 .hwp + 외부에서 구한 실제 샘플
    dst = os.path.join(OUT, 'hwp'); os.makedirs(dst, exist_ok=True)
    for root in (REAL_HWP, os.path.join(os.path.dirname(__file__), '..', 'test-files')):
        if not os.path.isdir(root): continue
        for dp, _, fs in os.walk(root):
            for f in fs:
                if f.lower().endswith('.hwp'):
                    shutil.copy(os.path.join(dp, f), os.path.join(dst, os.path.relpath(os.path.join(dp, f), root).replace(os.sep, '_')))

if __name__ == '__main__':
    steps = [make_docx, make_pptx, make_xlsx, make_pdf, make_korean_pdf, make_protected, make_hwpx, copy_real_hwp]
    for s in steps:
        try: s(); print('ok', s.__name__)
        except Exception as e: print('FAIL', s.__name__, repr(e)); 
    print(sorted(os.listdir(OUT)))
